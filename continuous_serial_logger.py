"""Log the Arduino continuous-mode stream to CSV/XLSX and plot it live."""

from __future__ import annotations

import argparse
import csv
import math
import os
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import TextIO


BAUD_RATE = 9600
MAX_RECORDS = 120
CSV_HEADER = "Minute,VPHS,VMAG"


@dataclass(frozen=True)
class Measurement:
    minute: int
    pc_datetime: datetime
    vphs: float
    vmag: float


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Record the Mega2560 continuous mode for 120 minutes."
    )
    parser.add_argument("--port", help="Serial port, for example COM4")
    parser.add_argument("--baud", type=int, default=BAUD_RATE)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path(__file__).resolve().parent / "measurements",
    )
    parser.add_argument("--no-plot", action="store_true", help="Disable live charts")
    parser.add_argument(
        "--simulate",
        action="store_true",
        help="Generate 120 test records without an Arduino",
    )
    parser.add_argument(
        "--simulation-interval",
        type=float,
        default=0.05,
        help="Seconds between simulated records",
    )
    return parser.parse_args()


def choose_serial_port(requested_port: str | None) -> str:
    if requested_port:
        return requested_port

    try:
        from serial.tools import list_ports
    except ImportError as exc:
        raise RuntimeError(
            "pyserial is not installed. Run: py -m pip install -r requirements.txt"
        ) from exc

    ports = sorted(list(list_ports.comports()), key=lambda item: item.device)
    if not ports:
        raise RuntimeError("No serial ports were found. Connect the Arduino and try again.")
    if len(ports) == 1:
        print(f"Using {ports[0].device}: {ports[0].description}")
        return ports[0].device

    print("Available serial ports:")
    for index, port in enumerate(ports, start=1):
        print(f"  {index}. {port.device} - {port.description}")

    while True:
        try:
            selection = int(input("Select port number: ").strip())
            if 1 <= selection <= len(ports):
                return ports[selection - 1].device
        except ValueError:
            pass
        print("Invalid selection. Please try again.")


class LivePlot:
    def __init__(self, enabled: bool) -> None:
        self.enabled = enabled
        self._plt = None
        self._figure = None
        self._axes = None
        self._lines = None

        if not enabled:
            return

        try:
            import matplotlib.pyplot as plt
        except ImportError as exc:
            raise RuntimeError(
                "matplotlib is not installed. Run: py -m pip install -r requirements.txt"
            ) from exc

        self._plt = plt
        plt.ion()
        self._figure, self._axes = plt.subplots(2, 1, figsize=(10, 7), sharex=True)
        self._figure.canvas.manager.set_window_title("Continuous Measurement - 2 Hours")
        self._lines = (
            self._axes[0].plot([], [], color="#2563EB", marker="o", markersize=3)[0],
            self._axes[1].plot([], [], color="#DC2626", marker="o", markersize=3)[0],
        )

        self._axes[0].set_title("VPHS - Real-time")
        self._axes[0].set_ylabel("VPHS (V)")
        self._axes[1].set_title("VMAG - Real-time")
        self._axes[1].set_ylabel("VMAG (V)")
        self._axes[1].set_xlabel("Elapsed time (minute)")
        for axis in self._axes:
            axis.set_xlim(1, MAX_RECORDS)
            axis.grid(True, alpha=0.3)
        self._figure.tight_layout()
        plt.show(block=False)

    def is_open(self) -> bool:
        if not self.enabled:
            return True
        return bool(self._plt.fignum_exists(self._figure.number))

    def update(self, records: list[Measurement]) -> None:
        if not self.enabled or not self.is_open():
            return

        minutes = [record.minute for record in records]
        vphs = [record.vphs for record in records]
        vmag = [record.vmag for record in records]
        self._lines[0].set_data(minutes, vphs)
        self._lines[1].set_data(minutes, vmag)
        for axis in self._axes:
            axis.relim()
            axis.autoscale_view(scalex=False, scaley=True)
        self._figure.canvas.draw_idle()
        self._plt.pause(0.01)

    def process_events(self) -> None:
        if self.enabled and self.is_open():
            self._plt.pause(0.01)

    def save_jpg(self, path: Path, records: list[Measurement]) -> None:
        if not records:
            return

        if self.enabled and self.is_open():
            self.update(records)
            self._figure.savefig(path, format="jpg", dpi=180, bbox_inches="tight")
            return

        try:
            import matplotlib.pyplot as plt
        except ImportError as exc:
            raise RuntimeError(
                "matplotlib is not installed. Run: py -m pip install -r requirements.txt"
            ) from exc

        minutes = [record.minute for record in records]
        figure, axes = plt.subplots(2, 1, figsize=(10, 7), sharex=True)
        axes[0].plot(minutes, [record.vphs for record in records], color="#2563EB")
        axes[0].set_title("VPHS - 2 Hour Measurement")
        axes[0].set_ylabel("VPHS (V)")
        axes[1].plot(minutes, [record.vmag for record in records], color="#DC2626")
        axes[1].set_title("VMAG - 2 Hour Measurement")
        axes[1].set_ylabel("VMAG (V)")
        axes[1].set_xlabel("Elapsed time (minute)")
        for axis in axes:
            axis.set_xlim(1, MAX_RECORDS)
            axis.grid(True, alpha=0.3)
        figure.tight_layout()
        figure.savefig(path, format="jpg", dpi=180, bbox_inches="tight")
        plt.close(figure)

    def close(self) -> None:
        if self.enabled and self.is_open():
            self._plt.ioff()
            self._plt.close(self._figure)


class SessionFiles:
    def __init__(self, output_dir: Path) -> None:
        self.started_at = datetime.now()
        timestamp = self.started_at.strftime("%Y%m%d_%H%M%S")
        output_dir.mkdir(parents=True, exist_ok=True)
        self.csv_path = output_dir / f"continuous_measurement_{timestamp}.csv"
        self.xlsx_path = output_dir / f"continuous_measurement_{timestamp}.xlsx"
        self.jpg_path = output_dir / f"continuous_measurement_{timestamp}.jpg"
        self._csv_file: TextIO = self.csv_path.open("w", newline="", encoding="utf-8-sig")
        self._csv_writer = csv.writer(self._csv_file)
        self._csv_writer.writerow(["Minute", "PC Date Time", "VPHS", "VMAG"])
        self._flush_csv()

    def append_csv(self, record: Measurement) -> None:
        self._csv_writer.writerow(
            [
                record.minute,
                record.pc_datetime.strftime("%Y-%m-%d %H:%M:%S"),
                f"{record.vphs:.3f}",
                f"{record.vmag:.3f}",
            ]
        )
        self._flush_csv()

    def _flush_csv(self) -> None:
        self._csv_file.flush()
        os.fsync(self._csv_file.fileno())

    def close(self) -> None:
        if not self._csv_file.closed:
            self._csv_file.close()


def save_excel(
    path: Path,
    records: list[Measurement],
    status: str,
    started_at: datetime,
) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.chart import LineChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is not installed. Run: py -m pip install -r requirements.txt"
        ) from exc

    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "Data"
    chart_sheet = workbook.create_sheet("Charts")
    data_sheet.sheet_view.showGridLines = False
    chart_sheet.sheet_view.showGridLines = False

    data_sheet["A1"] = "Session Status"
    data_sheet["B1"] = status
    data_sheet["A2"] = "Started At"
    data_sheet["B2"] = started_at
    data_sheet["A3"] = "Last Updated"
    data_sheet["B3"] = datetime.now()
    data_sheet["A4"] = "Records"
    data_sheet["B4"] = len(records)
    data_sheet.append(["Minute", "PC Date Time", "VPHS", "VMAG"])

    for record in records:
        data_sheet.append([record.minute, record.pc_datetime, record.vphs, record.vmag])

    title_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in data_sheet[5]:
        cell.fill = title_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for cell in data_sheet["A"][:4]:
        cell.font = Font(bold=True)

    data_sheet.freeze_panes = "A6"
    data_sheet.column_dimensions["A"].width = 12
    data_sheet.column_dimensions["B"].width = 22
    data_sheet.column_dimensions["C"].width = 14
    data_sheet.column_dimensions["D"].width = 14
    data_sheet["B2"].number_format = "yyyy-mm-dd hh:mm:ss"
    data_sheet["B3"].number_format = "yyyy-mm-dd hh:mm:ss"
    for row in range(6, 6 + len(records)):
        data_sheet.cell(row, 2).number_format = "yyyy-mm-dd hh:mm:ss"
        data_sheet.cell(row, 3).number_format = "0.000"
        data_sheet.cell(row, 4).number_format = "0.000"

    if records:
        last_row = 5 + len(records)
        table = Table(displayName="MeasurementData", ref=f"A5:D{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        data_sheet.add_table(table)

        for column, title, color, anchor in (
            (3, "VPHS by Minute", "2563EB", "A1"),
            (4, "VMAG by Minute", "DC2626", "A19"),
        ):
            chart = LineChart()
            chart.title = title
            chart.style = 13
            chart.height = 8
            chart.width = 16
            chart.y_axis.title = "Voltage (V)"
            chart.x_axis.title = "Elapsed time (minute)"
            values = Reference(data_sheet, min_col=column, min_row=5, max_row=last_row)
            categories = Reference(data_sheet, min_col=1, min_row=6, max_row=last_row)
            chart.add_data(values, titles_from_data=True)
            chart.set_categories(categories)
            chart.legend = None
            chart.series[0].graphicalProperties.line.solidFill = color
            chart.series[0].graphicalProperties.line.width = 22000
            chart_sheet.add_chart(chart, anchor)

    temporary_path = path.with_name(f".{path.stem}.tmp.xlsx")
    workbook.save(temporary_path)
    try:
        os.replace(temporary_path, path)
        return True
    except PermissionError:
        print(
            f"WARNING: Excel could not be updated because {path.name} is open. "
            "Close it; the logger will retry on the next record.",
            file=sys.stderr,
        )
        return False


def parse_measurement(line: str) -> tuple[int, float, float] | None:
    parts = [part.strip() for part in line.split(",")]
    if len(parts) != 3:
        return None
    try:
        minute = int(parts[0])
        vphs = float(parts[1])
        vmag = float(parts[2])
    except ValueError:
        return None
    if not 1 <= minute <= MAX_RECORDS:
        return None
    return minute, vphs, vmag


def add_record(
    values: tuple[int, float, float],
    records: list[Measurement],
    files: SessionFiles,
    plot: LivePlot,
) -> None:
    minute, vphs, vmag = values
    if any(record.minute == minute for record in records):
        print(f"Ignoring duplicate minute {minute}.")
        return
    if records and minute != records[-1].minute + 1:
        print(f"WARNING: expected minute {records[-1].minute + 1}, received {minute}.")

    record = Measurement(minute, datetime.now(), vphs, vmag)
    records.append(record)
    files.append_csv(record)
    save_excel(files.xlsx_path, records, "IN PROGRESS", files.started_at)
    plot.update(records)
    print(f"Minute {minute:3d}/{MAX_RECORDS}: VPHS={vphs:.3f}, VMAG={vmag:.3f}")


def finish_session(
    files: SessionFiles | None,
    records: list[Measurement],
    status: str,
    plot: LivePlot,
) -> None:
    if files is None:
        return
    files.close()
    saved = save_excel(files.xlsx_path, records, status, files.started_at)
    image_saved = False
    if status == "COMPLETED" and len(records) == MAX_RECORDS:
        plot.save_jpg(files.jpg_path, records)
        image_saved = True
    print(f"Session status: {status}; records: {len(records)}")
    print(f"CSV:  {files.csv_path}")
    if saved:
        print(f"Excel: {files.xlsx_path}")
    else:
        print("Close the Excel file and run the logger again to allow the final update.")
    if image_saved:
        print(f"Graph: {files.jpg_path}")


def run_simulation(args: argparse.Namespace, plot: LivePlot) -> int:
    print("Simulation mode: generating 120 records.")
    files = SessionFiles(args.output_dir)
    records: list[Measurement] = []
    status = "STOPPED EARLY"
    try:
        for minute in range(1, MAX_RECORDS + 1):
            if not plot.is_open():
                break
            vphs = 1.250 + 0.020 * math.sin(minute / 8.0)
            vmag = 2.500 + 0.015 * math.cos(minute / 10.0)
            add_record((minute, vphs, vmag), records, files, plot)
            if args.simulation_interval > 0:
                time.sleep(args.simulation_interval)
        if len(records) == MAX_RECORDS:
            status = "COMPLETED"
    except KeyboardInterrupt:
        print("Simulation stopped by user.")
    finally:
        finish_session(files, records, status, plot)
        plot.close()
    return 0


def run_serial(args: argparse.Namespace, plot: LivePlot) -> int:
    try:
        import serial
    except ImportError as exc:
        raise RuntimeError(
            "pyserial is not installed. Run: py -m pip install -r requirements.txt"
        ) from exc

    port = choose_serial_port(args.port)
    print(f"Opening {port} at {args.baud} baud...")
    print("Hold K1 for 3 seconds to enter Continuous mode.")

    files: SessionFiles | None = None
    records: list[Measurement] = []
    status = "STOPPED EARLY"
    serial_connection = None
    try:
        serial_connection = serial.Serial(port, args.baud, timeout=0.25)
        time.sleep(2.0)
        serial_connection.reset_input_buffer()

        while True:
            if not plot.is_open():
                print("Plot window closed by user.")
                break

            raw_line = serial_connection.readline()
            plot.process_events()
            if not raw_line:
                continue

            line = raw_line.decode("utf-8", errors="replace").strip()
            if line == CSV_HEADER:
                if files is not None:
                    finish_session(files, records, "RESTARTED BEFORE COMPLETION", plot)
                files = SessionFiles(args.output_dir)
                records = []
                print(f"Continuous session started. Output: {files.xlsx_path}")
                continue

            if line.startswith("#CONTINUOUS_COMPLETE"):
                if files is not None:
                    status = "COMPLETED" if len(records) == MAX_RECORDS else "INCOMPLETE"
                break

            if line.startswith("#CONTINUOUS_STOPPED"):
                print("Arduino left Continuous mode before the session completed.")
                break

            if files is None:
                continue

            values = parse_measurement(line)
            if values is None:
                continue
            add_record(values, records, files, plot)
            if len(records) >= MAX_RECORDS:
                status = "COMPLETED"
                break

    except KeyboardInterrupt:
        print("Stopped by user.")
    except serial.SerialException as exc:
        status = "SERIAL ERROR"
        print(f"Serial connection error: {exc}", file=sys.stderr)
    finally:
        if serial_connection is not None and serial_connection.is_open:
            serial_connection.close()
        finish_session(files, records, status, plot)
        plot.close()
    return 0 if status == "COMPLETED" else 1


def main() -> int:
    args = parse_args()
    plot = LivePlot(enabled=not args.no_plot)
    if args.simulate:
        return run_simulation(args, plot)
    return run_serial(args, plot)


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except RuntimeError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(2)
